"""Test setup. Real: FastAPI app, SQLAlchemy, Qdrant (embedded), storage, chunking, PDF/DOCX parsing, ownership checks.
Stand-ins (so tests run offline, fast, with no GPU): the embedding model, the LLM, and Firebase token verification.
"""
import hashlib
import io
import os
import re
import tempfile

import numpy as np
import pytest

_tmp = tempfile.mkdtemp(prefix="raggy-test-")
os.environ.update({
    "DATABASE_URL": f"sqlite:///{_tmp}/test.db", "STORAGE_BACKEND": "local", "LOCAL_STORAGE_DIR": f"{_tmp}/files",
    "QDRANT_URL": "", "QDRANT_PATH": f"{_tmp}/qdrant", "USE_RERANKER": "false", "RELEVANCE_THRESHOLD": "0.15",
    "WARMUP_EMBEDDER": "false", "FIREBASE_PROJECT_ID": "test-project", "ENVIRONMENT": "development",
})

from fastapi.testclient import TestClient  # noqa: E402

from app.services import firebase  # noqa: E402

USERS = {  # fake "ID tokens" -> verified claims
    "token-alice": {"uid": "uid-alice", "email": "alice@example.com", "name": "Alice"},
    "token-bob": {"uid": "uid-bob", "email": "bob@example.com", "name": "Bob"},
}


def _fake_verify(token):
    if token in USERS:
        return USERS[token]
    raise firebase.InvalidTokenError("bad token")


class FakeEmbedder:
    dim = 128

    def get_sentence_embedding_dimension(self):
        return self.dim

    def _vec(self, text):
        v = np.zeros(self.dim, dtype=np.float32)
        for w in re.findall(r"[a-z0-9]+", text.lower()):
            v[int(hashlib.md5(w.encode()).hexdigest(), 16) % self.dim] += 1.0
        n = np.linalg.norm(v)
        return v / n if n else v

    def encode(self, texts, **kw):
        if isinstance(texts, str):
            return self._vec(texts)
        return np.stack([self._vec(t) for t in texts])


class FakeLLM:
    def generate(self, messages):
        ctx = messages[-1]["content"].split("DOCUMENT CONTEXT:\n", 1)[1].split("\n", 2)[1]
        return f"Based on the document: {ctx[:80]}"


@pytest.fixture(scope="session", autouse=True)
def _patch_everything():
    mp = pytest.MonkeyPatch()
    mp.setattr(firebase, "verify_token", _fake_verify)
    from app.api import deps
    mp.setattr(deps, "verify_token", _fake_verify)
    from app.rag import embeddings, pipeline
    mp.setattr(embeddings, "get_embedding_model", lambda: FakeEmbedder())
    mp.setattr(pipeline, "get_llm_service", lambda: FakeLLM())
    yield
    mp.undo()


@pytest.fixture(scope="session")
def client():
    from app.main import app
    with TestClient(app) as c:
        yield c


def auth(user):
    return {"Authorization": f"Bearer token-{user}"}


def make_pdf(pages):
    import pymupdf
    d = pymupdf.open()
    for text in pages:
        p = d.new_page()
        p.insert_textbox(pymupdf.Rect(50, 50, 550, 780), text, fontsize=11)
    b = d.tobytes()
    d.close()
    return b


def make_docx(sections):
    import docx
    d = docx.Document()
    for heading, body in sections:
        d.add_heading(heading, level=1)
        d.add_paragraph(body)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


@pytest.fixture(scope="session")
def pdf_bytes():
    return make_pdf([
        "Introduction. This report studies solar panel efficiency in cold climates across northern regions.",
        "Methodology. We used a randomized controlled experiment with 120 rooftop installations over two winters.",
    ])


@pytest.fixture(scope="session")
def docx_bytes():
    return make_docx([
        ("Overview", "The quarterly financial report summarizes revenue growth and operating costs."),
        ("Methodology", "Interviews were conducted with forty customers using structured questionnaires."),
    ])
